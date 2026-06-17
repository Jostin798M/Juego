from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# ── Estilos globales ──────────────────────────────────────────────────────────
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

def set_col_width(table, col_idx, width_cm):
    for row in table.rows:
        row.cells[col_idx].width = Cm(width_cm)

def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def add_heading(doc, text, level=1, color='1F3864'):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        run.font.color.rgb = RGBColor.from_string(color)
        run.font.bold = True
    return p

def add_paragraph(doc, text, bold=False, italic=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    return p

def add_labeled(doc, label, value):
    p = doc.add_paragraph()
    r1 = p.add_run(label + ': ')
    r1.bold = True
    r1.font.size = Pt(11)
    r2 = p.add_run(value)
    r2.font.size = Pt(11)

def add_bullet(doc, text):
    p = doc.add_paragraph(text, style='List Bullet')
    p.runs[0].font.size = Pt(11)

def add_table_header(table, headers, fill='1F3864'):
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = h
        shade_cell(cell, fill)
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.size = Pt(10)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER

def fill_row(table, row_idx, values, center_cols=None):
    row = table.rows[row_idx]
    for i, val in enumerate(values):
        cell = row.cells[i]
        cell.text = val
        for para in cell.paragraphs:
            para.runs[0].font.size = Pt(10)
            if center_cols and i in center_cols:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER

# ══════════════════════════════════════════════════════════════════════════════
#  PORTADA
# ══════════════════════════════════════════════════════════════════════════════
doc.add_paragraph()
doc.add_paragraph()

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('BREWMANAGER')
run.bold = True
run.font.size = Pt(28)
run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run2 = subtitle.add_run('Sistema de Gestión para Bar-Café')
run2.font.size = Pt(16)
run2.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)

doc.add_paragraph()

curso = doc.add_paragraph()
curso.alignment = WD_ALIGN_PARAGRAPH.CENTER
curso.add_run('SOFG1006 – Desarrollo de Aplicaciones Web y Móviles\n').font.size = Pt(12)
curso.add_run('Escuela Superior Politécnica del Litoral – ESPOL').font.size = Pt(12)

doc.add_paragraph()

etapa_p = doc.add_paragraph()
etapa_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run3 = etapa_p.add_run('Documentación del Proyecto — Etapa 1')
run3.font.size = Pt(13)
run3.bold = True

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
#  ÍNDICE (manual)
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, 'Contenido', level=1)
items = [
    ('1.', 'Etapa 1: Análisis del sistema y levantamiento de requerimientos'),
    ('1.1', 'Condiciones académicas obligatorias (Prompt 1.0)'),
    ('1.2', 'Comprensión inicial del sistema (Prompt 1.1)'),
    ('1.3', 'Requerimientos funcionales y no funcionales (Prompt 1.6)'),
    ('1.4', 'Clasificación de requerimientos por fase (Prompt 1.7)'),
    ('1.5', 'Validación final de la etapa (Prompt 1.8)'),
]
for num, text in items:
    p = doc.add_paragraph()
    r = p.add_run(f'{num}  {text}')
    r.font.size = Pt(11)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
#  ETAPA 1
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, 'Etapa 1: Análisis del Sistema y Levantamiento de Requerimientos', level=1)
add_paragraph(doc,
    'Objetivo: comprender qué sistema web se va a desarrollar, qué problema resuelve, '
    'quién lo usará, qué procesos principales debe manejar y qué información será necesaria '
    'para las siguientes fases del proyecto.',
    italic=True)

doc.add_paragraph()

# ── 1.1 Condiciones obligatorias ─────────────────────────────────────────────
add_heading(doc, '1.1  Condiciones Académicas Obligatorias — Prompt 1.0', level=2)
add_paragraph(doc,
    'El proyecto BrewManager debe incluir obligatoriamente el módulo Clientes con los '
    'siguientes campos y reglas mínimas definidas por la cátedra.')

doc.add_paragraph()
add_paragraph(doc, 'Campos mínimos del módulo Clientes:', bold=True)

campos = [
    ('id', 'Identificador único generado automáticamente'),
    ('nombres', 'Obligatorio'),
    ('apellidos', 'Obligatorio'),
    ('identificacion', 'Obligatorio — debe ser único en el sistema'),
    ('telefono', 'Opcional'),
    ('celular', 'Obligatorio'),
    ('correo', 'Obligatorio — debe tener formato válido (ejemplo@dominio.com)'),
    ('direccion', 'Opcional'),
    ('estadoCivil', 'Obligatorio — valores: soltero, casado, divorciado, separado, unión libre'),
    ('estado', 'Obligatorio — valores: activo, inactivo'),
    ('fecha_registro', 'Se genera automáticamente al registrar el cliente'),
]

tbl = doc.add_table(rows=len(campos)+1, cols=2)
tbl.style = 'Table Grid'
add_table_header(tbl, ['Campo', 'Descripción / Regla'])
for i, (campo, desc) in enumerate(campos):
    fill_row(tbl, i+1, [campo, desc])

doc.add_paragraph()
add_paragraph(doc, 'Operaciones mínimas del módulo Clientes:', bold=True)
ops = [
    'Registrar clientes.',
    'Listar clientes.',
    'Buscar clientes por nombre, apellido o identificación.',
    'Consultar el detalle de un cliente.',
    'Editar la información de un cliente.',
    'Cambiar el estado civil del cliente.',
    'Cambiar el estado del cliente entre activo e inactivo.',
]
for op in ops:
    add_bullet(doc, op)

doc.add_paragraph()
add_paragraph(doc, 'Restricciones del proyecto en esta etapa:', bold=True)
restricciones = [
    'No se generará código en la etapa de análisis.',
    'No se diseñará la base de datos todavía.',
    'No se crearán APIs.',
    'No se desarrollará la aplicación móvil en esta etapa.',
    'No se deben crear funcionalidades no solicitadas.',
]
for r in restricciones:
    add_bullet(doc, r)

doc.add_page_break()

# ── 1.2 Comprensión inicial ──────────────────────────────────────────────────
add_heading(doc, '1.2  Comprensión Inicial del Sistema — Prompt 1.1', level=2)

add_paragraph(doc, 'Problema que resuelve', bold=True)
add_paragraph(doc,
    'Un bar-café generalmente lleva sus pedidos, ventas, inventario y clientes de forma '
    'manual o con herramientas dispersas (papel, cuadernos, hojas de cálculo). Esto genera '
    'errores en los pedidos, pérdida de información de clientes, dificultad para conocer el '
    'stock disponible y falta de control sobre las ventas del día.')
add_paragraph(doc,
    'BrewManager centraliza la gestión del bar-café en un sistema web accesible desde '
    'cualquier dispositivo, eliminando estos problemas.')

doc.add_paragraph()
add_paragraph(doc, 'Usuarios principales del sistema', bold=True)
usuarios = [
    ('Administrador', 'Gestiona todo el sistema: clientes, productos, pedidos, ventas y reportes.'),
    ('Cajero / Mesero', 'Registra pedidos, consulta el menú y atiende a los clientes.'),
]
tbl2 = doc.add_table(rows=len(usuarios)+1, cols=2)
tbl2.style = 'Table Grid'
add_table_header(tbl2, ['Rol', 'Responsabilidades'])
for i, (rol, resp) in enumerate(usuarios):
    fill_row(tbl2, i+1, [rol, resp])

doc.add_paragraph()
add_paragraph(doc, 'Procesos principales del negocio', bold=True)
procesos = [
    'Gestión de clientes — registrar, buscar, consultar, editar y cambiar estado.',
    'Gestión de productos/menú — registrar bebidas, cafés y snacks con precio y categoría.',
    'Gestión de pedidos — crear pedidos asociados a un cliente o mesa.',
    'Control de ventas — registrar y consultar ventas del día.',
    'Gestión de categorías — organizar el menú por tipo (bebidas calientes, frías, snacks, etc.).',
]
for p in procesos:
    add_bullet(doc, p)

doc.add_paragraph()
add_paragraph(doc, 'Fases de desarrollo del proyecto', bold=True)
fases = [
    'Fase 1: Desarrollo del frontend con datos simulados.',
    'Fase 2: Desarrollo del backend con Django y base de datos.',
    'Fase 3: Creación de APIs para conectar frontend con backend.',
    'Fase 4: Desarrollo de aplicación móvil con React Native consumiendo las mismas APIs.',
]
for f in fases:
    add_bullet(doc, f)

doc.add_paragraph()
add_paragraph(doc, 'Decisiones tomadas durante el análisis', bold=True)
decisiones = [
    'El sistema tendrá login con usuario y contraseña; solo usuarios registrados pueden acceder.',
    'Un pedido puede asociarse a un cliente registrado o indicar solo el número de mesa.',
    'Los estados de un pedido son: pendiente, en preparación, entregado, cancelado.',
    'Los productos manejan disponibilidad (disponible / no disponible), sin control de stock por ahora.',
    'Tanto el administrador como el cajero pueden registrar clientes.',
]
for d in decisiones:
    add_bullet(doc, d)

doc.add_page_break()

# ── 1.3 Requerimientos ───────────────────────────────────────────────────────
add_heading(doc, '1.3  Requerimientos Funcionales y No Funcionales — Prompt 1.6', level=2)

add_paragraph(doc, 'Requerimientos Funcionales', bold=True)
rf = [
    ('RF01', 'El sistema debe permitir iniciar sesión con usuario y contraseña.'),
    ('RF02', 'El sistema debe permitir registrar, listar, buscar, editar y consultar clientes.'),
    ('RF03', 'El sistema debe permitir cambiar el estado de un cliente (activo / inactivo).'),
    ('RF04', 'El sistema debe permitir cambiar el estado civil de un cliente.'),
    ('RF05', 'El sistema debe permitir registrar, listar, editar y cambiar la disponibilidad de productos.'),
    ('RF06', 'El sistema debe permitir organizar productos por categorías.'),
    ('RF07', 'El sistema debe permitir registrar un pedido asociado a un cliente o a una mesa.'),
    ('RF08', 'El sistema debe permitir agregar productos a un pedido.'),
    ('RF09', 'El sistema debe permitir cambiar el estado de un pedido (pendiente, en preparación, entregado, cancelado).'),
    ('RF10', 'El sistema debe mostrar el historial de pedidos.'),
    ('RF11', 'El sistema debe calcular el total de un pedido automáticamente.'),
    ('RF12', 'El sistema debe mostrar un resumen de ventas del día.'),
]
tbl3 = doc.add_table(rows=len(rf)+1, cols=2)
tbl3.style = 'Table Grid'
add_table_header(tbl3, ['Código', 'Descripción'])
for i, (cod, desc) in enumerate(rf):
    fill_row(tbl3, i+1, [cod, desc], center_cols=[0])

doc.add_paragraph()
add_paragraph(doc, 'Requerimientos No Funcionales', bold=True)
rnf = [
    ('RNF01', 'La interfaz debe ser responsive y funcionar en escritorio y dispositivos móviles.'),
    ('RNF02', 'El sistema debe validar los campos obligatorios en el navegador antes de enviar el formulario.'),
    ('RNF03', 'Los mensajes de error deben ser claros y comprensibles para el usuario.'),
    ('RNF04', 'La navegación debe ser intuitiva y accesible desde un menú principal visible.'),
    ('RNF05', 'El código debe estar comentado para facilitar su comprensión y revisión en clase.'),
    ('RNF06', 'El sistema debe usar datos simulados durante la fase de frontend.'),
    ('RNF07', 'La fase 1 se desarrollará con HTML, CSS, JavaScript y Bootstrap.'),
]
tbl4 = doc.add_table(rows=len(rnf)+1, cols=2)
tbl4.style = 'Table Grid'
add_table_header(tbl4, ['Código', 'Descripción'])
for i, (cod, desc) in enumerate(rnf):
    fill_row(tbl4, i+1, [cod, desc], center_cols=[0])

doc.add_page_break()

# ── 1.4 Clasificación por fase ───────────────────────────────────────────────
add_heading(doc, '1.4  Clasificación de Requerimientos por Fase — Prompt 1.7', level=2)

fases_data = [
    ('RF01', 'Pantalla de login visual', 'Lógica de autenticación real', 'Endpoint de autenticación', 'Pantalla de login móvil'),
    ('RF02', 'Formularios y tabla de clientes', 'Modelo Cliente + BD', 'GET / POST / PUT clientes', 'Pantallas de clientes móvil'),
    ('RF03', 'Botón cambiar estado (simulado)', 'Lógica de estado backend', 'PATCH estado cliente', 'Acción en app móvil'),
    ('RF04', 'Select estadoCivil en formulario', 'Validación en BD', 'Incluido en PUT cliente', 'Incluido en móvil'),
    ('RF05', 'Formularios y tabla de productos', 'Modelo Producto + BD', 'GET / POST / PUT productos', 'Pantallas de productos móvil'),
    ('RF06', 'Select de categoría en formulario', 'Modelo Categoría + BD', 'GET categorías', 'Incluido en móvil'),
    ('RF07', 'Formulario de pedido (simulado)', 'Modelo Pedido + BD', 'POST pedido', 'Pantalla de pedido móvil'),
    ('RF08', 'Selección visual de productos', 'Relación Pedido-Producto', 'POST detalle pedido', 'Incluido en móvil'),
    ('RF09', 'Botón cambiar estado pedido', 'Lógica de estados backend', 'PATCH estado pedido', 'Acción en app móvil'),
    ('RF10', 'Tabla historial (simulada)', 'Consulta BD historial', 'GET pedidos', 'Lista de pedidos móvil'),
    ('RF11', 'Cálculo de total en JavaScript', 'Calculado en backend', 'Incluido en respuesta GET', 'Incluido en móvil'),
    ('RF12', 'Sección resumen visual (simulada)', 'Consulta agregada BD', 'GET resumen ventas', 'Panel resumen móvil'),
]

tbl5 = doc.add_table(rows=len(fases_data)+1, cols=5)
tbl5.style = 'Table Grid'
add_table_header(tbl5, ['Req.', 'Fase 1\nFrontend', 'Fase 2\nBackend/BD', 'Fase 3\nAPIs', 'Fase 4\nMóvil'])
for i, row_data in enumerate(fases_data):
    fill_row(tbl5, i+1, list(row_data), center_cols=[0])

doc.add_page_break()

# ── 1.5 Validación final ─────────────────────────────────────────────────────
add_heading(doc, '1.5  Validación Final de la Etapa 1 — Prompt 1.8', level=2)

criterios = [
    ('El problema está claramente definido', '✓'),
    ('Los usuarios del sistema están identificados', '✓'),
    ('Los procesos principales están claros', '✓'),
    ('Los datos y campos principales están definidos', '✓'),
    ('Las reglas de negocio están claras', '✓'),
    ('Los requerimientos son verificables y están numerados', '✓'),
    ('Existe información suficiente para diseñar pantallas, formularios y navegación', '✓'),
]
tbl6 = doc.add_table(rows=len(criterios)+1, cols=2)
tbl6.style = 'Table Grid'
add_table_header(tbl6, ['Criterio de validación', 'Estado'])
for i, (crit, est) in enumerate(criterios):
    fill_row(tbl6, i+1, [crit, est], center_cols=[1])

doc.add_paragraph()
resultado = doc.add_paragraph()
resultado.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_r = resultado.add_run('RESULTADO: Aprobado para pasar al diseño del frontend.')
run_r.bold = True
run_r.font.size = Pt(12)
run_r.font.color.rgb = RGBColor(0x1F, 0x85, 0x3C)

# ══════════════════════════════════════════════════════════════════════════════
#  GUARDAR
# ══════════════════════════════════════════════════════════════════════════════
doc.save('/home/user/Juego/BrewManager_Documentacion.docx')
print('Documento generado correctamente.')
